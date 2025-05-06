from django.shortcuts import render, redirect, get_object_or_404
from .models import Orcamento, ItemOrcamento, JustificativaAprovacao
from produtos.models import Produto, Fornecedor
from pedidos.models import Pedido, ItemPedido
from escolas.models import Escola
from django.contrib import messages
from django.utils import timezone
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse, JsonResponse
from django.contrib.auth.decorators import login_required, permission_required
from django.db import transaction
from django.urls import reverse
from django.template.loader import render_to_string
import weasyprint
from django.conf import settings
import os
from openpyxl.utils import get_column_letter

def lista(request):
    orcamentos = Orcamento.objects.all().order_by('-data_criacao')
    return render(request, 'orcamentos/lista.html', {'orcamentos': orcamentos})

def novo(request):
    produtos = Produto.objects.filter(ativo=True).order_by('nome').values('id', 'nome', 'valor_unitario')
    fornecedores = Fornecedor.objects.filter(ativo=True).order_by('razao_social')
    if request.method == 'POST':
        nome = request.POST.get('nome')
        forma_pagamento = request.POST.get('forma_pagamento')
        data_entrega_prevista = request.POST.get('data_entrega_prevista')
        itens = []
        for idx in range(0, 20):  # Suporta até 20 itens por orçamento
            produto_id = request.POST.get(f'produto_{idx}')
            quantidade = request.POST.get(f'quantidade_{idx}')
            valor_unitario = request.POST.get(f'valor_unitario_{idx}')
            fornecedor_id = request.POST.get(f'fornecedor_{idx}')
            produto_manual = request.POST.get(f'produto_manual_{idx}')
            
            # Se não houver mais itens, sair do loop
            if not (produto_id or produto_manual) and not quantidade and not valor_unitario:
                break
                
            if quantidade and valor_unitario:
                try:
                    # Se for produto cadastrado
                    if produto_id:
                        produto = Produto.objects.get(id=produto_id)
                        nome_produto = produto.nome
                    # Se for produto manual
                    else:
                        nome_produto = produto_manual
                        produto = None
                        
                    fornecedor = Fornecedor.objects.get(id=fornecedor_id) if fornecedor_id else None
                    
                    itens.append({
                        'produto': produto,
                        'nome_produto': nome_produto,
                        'quantidade': int(quantidade),
                        'valor_unitario': float(valor_unitario),
                        'fornecedor': fornecedor
                    })
                except (Produto.DoesNotExist, Fornecedor.DoesNotExist):
                    continue
        if not nome or not itens:
            messages.error(request, 'Preencha o nome do orçamento e adicione pelo menos um item.')
            return render(request, 'orcamentos/novo.html', {'produtos': produtos, 'fornecedores': fornecedores})
        orcamento = Orcamento.objects.create(
            nome=nome, 
            data_criacao=timezone.now(),
            forma_pagamento=forma_pagamento,
            data_entrega_prevista=data_entrega_prevista if data_entrega_prevista else None
        )
        for item in itens:
            ItemOrcamento.objects.create(
                orcamento=orcamento,
                produto=item['produto'],
                quantidade=item['quantidade'],
                valor_unitario=item['valor_unitario'],
                fornecedor=item['fornecedor'],
                total=item['quantidade'] * item['valor_unitario']
            )
        messages.success(request, 'Orçamento cadastrado com sucesso!')
        return redirect('orcamentos:lista')
    return render(request, 'orcamentos/novo.html', {'produtos': produtos, 'fornecedores': fornecedores})

def importar(request):
    produtos = Produto.objects.filter(ativo=True).order_by('nome')
    fornecedores = Fornecedor.objects.filter(ativo=True).order_by('razao_social')
    if request.method == 'POST' and 'arquivo' in request.FILES:
        arquivo = request.FILES['arquivo']
        if not arquivo.name.endswith(('.xlsx', '.xls')):
            messages.error(request, 'Formato de arquivo não suportado. Envie um arquivo Excel (.xlsx ou .xls).')
            return redirect('orcamentos:importar')
        try:
            df = pd.read_excel(arquivo)
            colunas_necessarias = ['Nome do Orçamento', 'Produto', 'Quantidade', 'Valor Unitário', 'Fornecedor (CNPJ)']
            for coluna in colunas_necessarias:
                if coluna not in df.columns:
                    messages.error(request, f'Coluna "{coluna}" não encontrada no arquivo.')
                    return redirect('orcamentos:importar')
            orcamentos_criados = 0
            erros = []
            for nome_orcamento, grupo in df.groupby('Nome do Orçamento'):
                forma_pagamento = 'avista'  # valor padrão
                data_entrega_prevista = None
                
                # Verificar se existem as colunas opcionais
                if 'Forma de Pagamento' in df.columns and not grupo['Forma de Pagamento'].isnull().all():
                    forma_pagamento_valor = grupo['Forma de Pagamento'].iloc[0]
                    # Mapear valores descritivos para os códigos internos
                    mapeamento_pagamento = {
                        'À Vista': 'avista',
                        'A Vista': 'avista',
                        '30 Dias': '30dias',
                        '60 Dias': '60dias',
                        '90 Dias': '90dias',
                        'Parcelado': 'parcelado',
                        'Outros': 'outros'
                    }
                    forma_pagamento = mapeamento_pagamento.get(forma_pagamento_valor, 'avista')
                
                if 'Data Prevista de Entrega' in df.columns and not grupo['Data Prevista de Entrega'].isnull().all():
                    data_entrega_str = grupo['Data Prevista de Entrega'].iloc[0]
                    if pd.notna(data_entrega_str):
                        try:
                            if isinstance(data_entrega_str, str):
                                # Tentar converter string para data
                                from datetime import datetime
                                data_entrega_prevista = datetime.strptime(data_entrega_str, '%d/%m/%Y').date()
                            else:
                                # Se já for um objeto datetime pandas
                                data_entrega_prevista = data_entrega_str
                        except:
                            data_entrega_prevista = None
                
                orcamento = Orcamento.objects.create(
                    nome=nome_orcamento, 
                    forma_pagamento=forma_pagamento,
                    data_entrega_prevista=data_entrega_prevista
                )
                for _, row in grupo.iterrows():
                    try:
                        produto = Produto.objects.get(nome__iexact=row['Produto'])
                        fornecedor = None
                        if pd.notna(row['Fornecedor (CNPJ)']):
                            fornecedor = Fornecedor.objects.filter(cnpj=str(row['Fornecedor (CNPJ)']).strip()).first()
                        ItemOrcamento.objects.create(
                            orcamento=orcamento,
                            produto=produto,
                            quantidade=int(row['Quantidade']),
                            valor_unitario=float(row['Valor Unitário']),
                            fornecedor=fornecedor,
                            total=int(row['Quantidade']) * float(row['Valor Unitário'])
                        )
                    except Produto.DoesNotExist:
                        erros.append(f"Produto '{row['Produto']}' não encontrado no sistema.")
                orcamentos_criados += 1
            if orcamentos_criados:
                messages.success(request, f'{orcamentos_criados} orçamento(s) importado(s) com sucesso!')
            if erros:
                messages.warning(request, 'Ocorreram erros em alguns itens: ' + '; '.join(erros[:5]) + ('. E mais...' if len(erros) > 5 else ''))
            return redirect('orcamentos:lista')
        except Exception as e:
            messages.error(request, f'Erro ao processar o arquivo: {str(e)}')
            return redirect('orcamentos:importar')
    return render(request, 'orcamentos/importar.html', {'produtos': produtos, 'fornecedores': fornecedores})

def download_modelo(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Modelo Orçamentos"
    cabecalhos = ["Nome do Orçamento", "Produto", "Quantidade", "Valor Unitário", "Fornecedor (CNPJ)", "Forma de Pagamento", "Data Prevista de Entrega"]
    for col_num, cabecalho in enumerate(cabecalhos, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = cabecalho
    exemplos = [
        ["Orçamento Maio/2025", "Caderno Espiral", 10, 15.90, "12345678000199", "À Vista", "31/05/2025"],
        ["Orçamento Maio/2025", "Lápis HB", 20, 2.50, "12345678000199", "À Vista", "31/05/2025"],
        ["Orçamento Junho/2025", "Caderno Espiral", 5, 16.00, "", "30 Dias", "30/06/2025"],
    ]
    for row_num, exemplo in enumerate(exemplos, 2):
        for col_num, valor in enumerate(exemplo, 1):
            sheet.cell(row=row_num, column=col_num).value = valor
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=modelo_orcamentos.xlsx'
    workbook.save(response)
    return response

@login_required
def detalhe(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    itens = orcamento.itens.select_related('produto', 'fornecedor').all()
    return render(request, 'orcamentos/detalhe.html', {
        'orcamento': orcamento, 
        'itens': itens
    })

@login_required
def aprovar(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    if request.method == 'POST':
        justificativa = request.POST.get('justificativa')
        orcamento.status = 'aprovado'
        orcamento.justificativa = justificativa
        orcamento.data_aprovacao = timezone.now()
        orcamento.usuario_aprovador = request.user
        orcamento.save()
        messages.success(request, 'Orçamento aprovado com sucesso!')
        return redirect('orcamentos:detalhe', orcamento_id=orcamento.id)
    return render(request, 'orcamentos/aprovar.html', {'orcamento': orcamento})

@login_required
def rejeitar(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    if request.method == 'POST':
        justificativa = request.POST.get('justificativa')
        orcamento.status = 'rejeitado'
        orcamento.justificativa = justificativa
        orcamento.data_aprovacao = timezone.now()
        orcamento.usuario_aprovador = request.user
        orcamento.save()
        messages.success(request, 'Orçamento rejeitado!')
        return redirect('orcamentos:detalhe', orcamento_id=orcamento.id)
    return render(request, 'orcamentos/rejeitar.html', {'orcamento': orcamento})

def novo_orcamento(request):
    if request.method == 'POST':
        form = OrcamentoForm(request.POST)
        if form.is_valid():
            orcamento = form.save(commit=False)
            orcamento.criado_por = request.user
            orcamento.save()

            # Processar itens do orçamento
            i = 1
            while True:
                produto_id = request.POST.get(f'produto_{i}')
                produto_manual = request.POST.get(f'produto_manual_{i}')
                quantidade = request.POST.get(f'quantidade_{i}')
                valor_unitario = request.POST.get(f'valor_unitario_{i}')

                if not (produto_id or produto_manual) or not quantidade or not valor_unitario:
                    break

                if produto_id:
                    produto = Produto.objects.get(id=produto_id)
                else:
                    produto = None
                    descricao_produto = produto_manual

                ItemOrcamento.objects.create(
                    orcamento=orcamento,
                    produto=produto,
                    descricao_produto=descricao_produto if not produto else None,
                    quantidade=quantidade,
                    valor_unitario=valor_unitario
                )
                i += 1

            messages.success(request, 'Orçamento criado com sucesso!')
            return redirect('orcamentos:listar')
    else:
        form = OrcamentoForm()

    produtos = Produto.objects.all()
    return render(request, 'orcamentos/novo.html', {
        'form': form,
        'produtos': produtos
    })

def editar(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    if orcamento.status != 'pendente':
        messages.error(request, 'Só é possível editar orçamentos pendentes.')
        return redirect('orcamentos:detalhe', orcamento_id=orcamento.id)
    produtos = Produto.objects.filter(ativo=True).order_by('nome').values('id', 'nome', 'valor_unitario')
    fornecedores = Fornecedor.objects.filter(ativo=True).order_by('razao_social')
    itens_existentes = orcamento.itens.all()
    if request.method == 'POST':
        nome = request.POST.get('nome')
        forma_pagamento = request.POST.get('forma_pagamento')
        data_entrega_prevista = request.POST.get('data_entrega_prevista')
        itens = []
        for idx in range(0, 20):
            produto_id = request.POST.get(f'produto_{idx}')
            quantidade = request.POST.get(f'quantidade_{idx}')
            valor_unitario = request.POST.get(f'valor_unitario_{idx}')
            fornecedor_id = request.POST.get(f'fornecedor_{idx}')
            produto_manual = request.POST.get(f'produto_manual_{idx}')
            if not (produto_id or produto_manual) and not quantidade and not valor_unitario:
                break
            if quantidade and valor_unitario:
                try:
                    if produto_id:
                        produto = Produto.objects.get(id=produto_id)
                        nome_produto = produto.nome
                    else:
                        nome_produto = produto_manual
                        produto = None
                    fornecedor = Fornecedor.objects.get(id=fornecedor_id) if fornecedor_id else None
                    itens.append({
                        'produto': produto,
                        'nome_produto': nome_produto,
                        'quantidade': int(quantidade),
                        'valor_unitario': float(valor_unitario),
                        'fornecedor': fornecedor
                    })
                except (Produto.DoesNotExist, Fornecedor.DoesNotExist):
                    continue
        if not nome or not itens:
            messages.error(request, 'Preencha o nome do orçamento e adicione pelo menos um item.')
            return render(request, 'orcamentos/editar.html', {'produtos': produtos, 'fornecedores': fornecedores, 'orcamento': orcamento, 'itens_existentes': itens_existentes})
        orcamento.nome = nome
        orcamento.forma_pagamento = forma_pagamento
        orcamento.data_entrega_prevista = data_entrega_prevista if data_entrega_prevista else None
        orcamento.save()
        orcamento.itens.all().delete()
        for item in itens:
            ItemOrcamento.objects.create(
                orcamento=orcamento,
                produto=item['produto'],
                quantidade=item['quantidade'],
                valor_unitario=item['valor_unitario'],
                fornecedor=item['fornecedor'],
                total=item['quantidade'] * item['valor_unitario']
            )
        messages.success(request, 'Orçamento atualizado com sucesso!')
        return redirect('orcamentos:detalhe', orcamento_id=orcamento.id)
    return render(request, 'orcamentos/editar.html', {'produtos': produtos, 'fornecedores': fornecedores, 'orcamento': orcamento, 'itens_existentes': itens_existentes})

def excluir(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    if orcamento.status != 'pendente':
        messages.error(request, 'Não é possível excluir orçamentos que já foram aprovados ou rejeitados.')
        return redirect('orcamentos:lista')
    
    nome_orcamento = orcamento.nome
    orcamento.delete()
    messages.success(request, f'Orçamento "{nome_orcamento}" excluído com sucesso!')
    return redirect('orcamentos:lista')

def comparar(request):
    # Processar parâmetros de busca
    nome = request.GET.get('nome', '')
    status = request.GET.get('status', '')
    data_inicio = request.GET.get('data_inicio', '')
    
    # Query base para busca
    orcamentos_query = Orcamento.objects.all()
    
    # Aplicar filtros
    if nome:
        orcamentos_query = orcamentos_query.filter(nome__icontains=nome)
    if status:
        orcamentos_query = orcamentos_query.filter(status=status)
    if data_inicio:
        orcamentos_query = orcamentos_query.filter(data_criacao__date__gte=data_inicio)
    
    # Ordenar por data de criação decrescente
    orcamentos_encontrados = orcamentos_query.order_by('-data_criacao')
    
    # Obtém os nomes únicos de orçamentos para agrupamento
    nomes_orcamentos = Orcamento.objects.values_list('nome', flat=True).distinct()
    
    # Dicionário que armazenará orçamentos agrupados por nome
    grupos_orcamentos = {}
    
    # Para cada nome de orçamento, agrupa os orçamentos relacionados
    for nome in nomes_orcamentos:
        orcamentos = Orcamento.objects.filter(nome=nome).order_by('id')
        if orcamentos.count() > 0:
            grupos_orcamentos[nome] = orcamentos
    
    # Para cada grupo, calcular qual tem o menor valor total
    grupos_com_analise = []
    for nome, orcamentos in grupos_orcamentos.items():
        # Se houver apenas 1 orçamento no grupo, não há necessidade de comparação
        if len(orcamentos) < 2:
            continue
            
        # Encontrar o orçamento com menor valor total
        menor_valor = float('inf')
        orcamento_menor_valor = None
        
        for orcamento in orcamentos:
            valor_total = orcamento.total
            if valor_total < menor_valor:
                menor_valor = valor_total
                orcamento_menor_valor = orcamento
        
        # Calcular a economia em relação ao segundo menor
        segundo_menor = float('inf')
        for orcamento in orcamentos:
            if orcamento.id != orcamento_menor_valor.id and orcamento.total < segundo_menor:
                segundo_menor = orcamento.total
        
        economia = segundo_menor - menor_valor if segundo_menor != float('inf') else 0
        economia_percentual = (economia / segundo_menor) * 100 if segundo_menor != float('inf') else 0
        
        # Coletar todos os produtos únicos do grupo e preparar comparação de itens
        produtos_unicos = set()
        itens_comparacao = []
        
        # Primeiro, colete todos os produtos únicos
        for orcamento in orcamentos:
            for item in orcamento.itens.all():
                if item.produto:
                    produtos_unicos.add(item.produto.nome)
                else:
                    produtos_unicos.add(item.descricao_produto)
        
        # Para cada produto único, compare os preços entre os orçamentos
        for produto in sorted(produtos_unicos):
            precos = []
            menor_preco = float('inf')
            
            # Coletar preços de cada orçamento para este produto
            for orcamento in orcamentos:
                preco = None
                for item in orcamento.itens.all():
                    nome_produto = item.produto.nome if item.produto else item.descricao_produto
                    if nome_produto == produto:
                        preco = item.valor_unitario
                        if preco < menor_preco:
                            menor_preco = preco
                        break
                precos.append({
                    'valor': preco if preco is not None else 0,
                    'melhor_preco': False  # Será atualizado depois
                })
            
            # Marcar o(s) melhor(es) preço(s)
            for preco in precos:
                if preco['valor'] == menor_preco:
                    preco['melhor_preco'] = True
            
            itens_comparacao.append({
                'descricao': produto,
                'precos': precos
            })
        
        # Adicionar ao resultado
        grupos_com_analise.append({
            'nome': nome,
            'orcamentos': orcamentos,
            'orcamento_melhor': orcamento_menor_valor,
            'economia': economia,
            'economia_percentual': economia_percentual,
            'quantidade_orcamentos': len(orcamentos),
            'produtos_unicos': sorted(list(produtos_unicos)),
            'itens_comparacao': itens_comparacao  # Adicionando a comparação de itens
        })
    
    return render(request, 'orcamentos/comparar.html', {
        'grupos': grupos_com_analise,
        'orcamentos_encontrados': orcamentos_encontrados
    })

def definir_melhor(request):
    if request.method == 'POST':
        orcamento_id = request.POST.get('orcamento_id')
        if orcamento_id:
            try:
                # Primeiro, remove a marcação de todos os orçamentos com o mesmo nome
                orcamento = get_object_or_404(Orcamento, id=orcamento_id)
                Orcamento.objects.filter(nome=orcamento.nome).update(melhor_orcamento=False)
                
                # Define este como o melhor
                orcamento.melhor_orcamento = True
                orcamento.save()
                
                messages.success(request, f'Orçamento "{orcamento.nome}" (#{orcamento.id}) definido como melhor opção.')
            except Exception as e:
                messages.error(request, f'Erro ao definir melhor orçamento: {str(e)}')
    
    return redirect('orcamentos:comparar')

def agrupar_orcamentos(request):
    if request.method == 'POST':
        orcamento_ids = request.POST.getlist('orcamentos[]')
        if orcamento_ids and len(orcamento_ids) >= 2:
            try:
                # Obter os orçamentos selecionados
                orcamentos = Orcamento.objects.filter(id__in=orcamento_ids)
                
                # Usar o nome do primeiro orçamento como padrão para o grupo
                nome_padrao = orcamentos.first().nome
                
                # Atualizar todos os orçamentos com o mesmo nome
                for orcamento in orcamentos:
                    if orcamento.nome != nome_padrao:
                        orcamento.nome = nome_padrao
                        orcamento.save()
                
                messages.success(request, f'{len(orcamentos)} orçamentos foram agrupados com sucesso sob o nome "{nome_padrao}".')
            except Exception as e:
                messages.error(request, f'Erro ao agrupar orçamentos: {str(e)}')
        else:
            messages.warning(request, 'Selecione ao menos 2 orçamentos para agrupar.')
    
    return redirect('orcamentos:comparar')

@login_required
@transaction.atomic
def aprovar_orcamento(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    
    if request.method == 'POST':
        # Verificar se é o melhor preço do grupo
        grupo_orcamentos = Orcamento.objects.filter(nome=orcamento.nome).exclude(id=orcamento.id)
        melhor_preco = min(grupo_orcamentos, key=lambda x: x.total, default=None)
        
        # Se não for o melhor preço, verificar se tem justificativa
        if melhor_preco and orcamento.total > melhor_preco.total:
            justificativa = request.POST.get('justificativa')
            if not justificativa:
                messages.error(request, 'É necessário fornecer uma justificativa para aprovar um orçamento que não tem o melhor preço.')
                return JsonResponse({
                    'status': 'error',
                    'message': 'Justificativa necessária',
                    'requires_justification': True,
                    'redirect_url': None
                })
            
            # Criar justificativa
            JustificativaAprovacao.objects.create(
                orcamento=orcamento,
                texto=justificativa,
                usuario=request.user
            )
        
        # Aprovar o orçamento selecionado
        orcamento.status = 'aprovado'
        orcamento.save()
        
        # Rejeitar os outros orçamentos do grupo
        grupo_orcamentos.update(status='rejeitado')
        
        messages.success(request, f'Orçamento #{orcamento.id} aprovado com sucesso!')
        
        return JsonResponse({
            'status': 'success',
            'message': 'Orçamento aprovado com sucesso!',
            'requires_justification': False,
            'redirect_url': reverse('orcamentos:comparar')
        })
    
    return JsonResponse({
        'status': 'error',
        'message': 'Método não permitido',
        'requires_justification': False,
        'redirect_url': None
    })

@login_required
def exportar_pdf(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    itens = orcamento.itens.select_related('produto', 'fornecedor').all()
    
    # Renderizar o template HTML
    html_string = render_to_string('orcamentos/pdf_template.html', {
        'orcamento': orcamento,
        'itens': itens,
        'STATIC_URL': settings.STATIC_URL
    })
    
    # Criar o arquivo PDF
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="orcamento_{orcamento.id}.pdf"'
    
    # Gerar PDF usando WeasyPrint
    pdf = weasyprint.HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
    response.write(pdf)
    
    return response

@login_required
def exportar_excel(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    itens = orcamento.itens.select_related('produto', 'fornecedor').all()
    
    # Criar um novo workbook e selecionar a planilha ativa
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Orçamento #{orcamento.id}"
    
    # Estilo para cabeçalhos
    header_style = Font(bold=True)
    header_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # Informações do orçamento
    sheet['A1'] = "INFORMAÇÕES DO ORÇAMENTO"
    sheet['A1'].font = Font(bold=True, size=14)
    sheet.merge_cells('A1:E1')
    
    info_rows = [
        ("Nome:", orcamento.nome),
        ("Status:", orcamento.get_status_display()),
        ("Data de Criação:", orcamento.data_criacao.strftime("%d/%m/%Y %H:%M")),
        ("Forma de Pagamento:", orcamento.get_forma_pagamento_display() if orcamento.forma_pagamento else "Não informada"),
        ("Data Prevista de Entrega:", orcamento.data_entrega_prevista.strftime("%d/%m/%Y") if orcamento.data_entrega_prevista else "Não informada"),
    ]
    
    current_row = 2
    for label, value in info_rows:
        sheet[f'A{current_row}'] = label
        sheet[f'B{current_row}'] = value
        sheet[f'A{current_row}'].font = Font(bold=True)
        current_row += 1
    
    # Espaço entre informações e itens
    current_row += 1
    
    # Cabeçalho dos itens
    headers = ["Produto", "Quantidade", "Valor Unitário", "Total", "Fornecedor"]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=current_row, column=col)
        cell.value = header
        cell.font = header_style
        cell.fill = header_fill
    
    # Dados dos itens
    for item in itens:
        current_row += 1
        row_data = [
            item.produto.nome if item.produto else item.descricao_produto,
            item.quantidade,
            float(item.valor_unitario),
            float(item.total),
            item.fornecedor.razao_social if item.fornecedor else "-"
        ]
        for col, value in enumerate(row_data, 1):
            cell = sheet.cell(row=current_row, column=col)
            cell.value = value
            if col in [3, 4]:  # Formatar valores monetários
                cell.number_format = 'R$ #,##0.00'
    
    # Total
    current_row += 1
    sheet.cell(row=current_row, column=3, value="Total:").font = Font(bold=True)
    total_cell = sheet.cell(row=current_row, column=4, value=float(orcamento.total))
    total_cell.font = Font(bold=True)
    total_cell.number_format = 'R$ #,##0.00'
    
    # Ajustar largura das colunas
    for col in range(1, len(headers) + 1):
        max_length = 0
        for row in range(1, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=col)
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[get_column_letter(col)].width = adjusted_width
    
    # Criar a resposta HTTP
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=orcamento_{orcamento.id}.xlsx'
    
    # Salvar o arquivo
    workbook.save(response)
    return response

@login_required
@permission_required('pedidos.add_pedido')
def gerar_pedido(request, orcamento_id):
    orcamento = get_object_or_404(Orcamento, id=orcamento_id)
    
    # Verificar se o orçamento está aprovado
    if orcamento.status != 'aprovado':
        messages.error(request, 'Só é possível gerar pedido a partir de orçamentos aprovados.')
        return redirect('orcamentos:detalhe', orcamento_id=orcamento.id)
    
    if request.method == 'POST':
        escola_id = request.POST.get('escola')
        
        if not escola_id:
            messages.error(request, 'Selecione uma escola para o pedido.')
            escolas = Escola.objects.filter(ativo=True).order_by('nome')
            return render(request, 'orcamentos/gerar_pedido.html', {
                'orcamento': orcamento,
                'escolas': escolas
            })
        
        try:
            with transaction.atomic():
                escola = Escola.objects.get(id=escola_id)
                
                # Criar o pedido
                pedido = Pedido.objects.create(
                    escola=escola,
                    data_solicitacao=timezone.now(),
                    status='pendente',
                    observacoes=f'Pedido gerado a partir do orçamento #{orcamento.id}'
                )
                
                # Criar os itens do pedido
                for item_orcamento in orcamento.itens.all():
                    ItemPedido.objects.create(
                        pedido=pedido,
                        produto=item_orcamento.produto,
                        quantidade=item_orcamento.quantidade,
                        valor_unitario=item_orcamento.valor_unitario
                    )
                
                messages.success(request, f'Pedido #{pedido.id} gerado com sucesso!')
                return redirect('pedidos:detalhes', pk=pedido.id)
                
        except Escola.DoesNotExist:
            messages.error(request, 'Escola não encontrada.')
        except Exception as e:
            messages.error(request, f'Erro ao gerar pedido: {str(e)}')
    
    # Se for GET ou houver erro, mostrar formulário
    escolas = Escola.objects.filter(ativo=True).order_by('nome')
    return render(request, 'orcamentos/gerar_pedido.html', {
        'orcamento': orcamento,
        'escolas': escolas
    })
