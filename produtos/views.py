from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from .models import Produto, Fornecedor
import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.db import transaction
from decimal import Decimal
from django.utils.decorators import method_decorator
from django.contrib.auth.decorators import login_required
from django.db.models import Q

def get_client_ip(request):
    """Obtém o endereço IP do cliente"""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0]
    else:
        ip = request.META.get('REMOTE_ADDR')
    return ip

def lista(request):
    """Exibe a lista de produtos"""
    produtos = Produto.objects.all()
    return render(request, 'produtos/lista.html', {'produtos': produtos})

def novo(request):
    """Adiciona um novo produto"""
    if request.method == 'POST':
        # Extrai os dados do formulário
        nome = request.POST.get('nome')
        descricao = request.POST.get('descricao')
        valor_unitario = request.POST.get('valor_unitario').replace(',', '.')
        unidade_medida = request.POST.get('unidade_medida')
        codigo = request.POST.get('codigo')
        fornecedor_id = request.POST.get('fornecedor')
        
        # Valida os dados
        if not nome:
            messages.error(request, 'O nome do produto é obrigatório.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': request.POST, 'fornecedores': fornecedores})
        
        # Remove espaços extras do nome para evitar duplicações por espaços
        nome = nome.strip()
        
        # Verificação de produto duplicado pelo nome (PRIORITÁRIA)
        nome_existente = Produto.objects.filter(nome__iexact=nome).first()
        if nome_existente:
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o nome "{nome}".')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': request.POST, 'fornecedores': fornecedores})
        
        try:
            valor_unitario = Decimal(valor_unitario)
            if valor_unitario <= 0:
                raise ValueError("Valor deve ser maior que zero")
        except:
            messages.error(request, 'O valor unitário deve ser um número positivo.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': request.POST, 'fornecedores': fornecedores})
            
        # Verificações secundárias de duplicação
        # 2. Verificar se já existe produto com o mesmo código (se o código foi fornecido)
        if codigo and Produto.objects.filter(codigo=codigo).exists():
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o código "{codigo}".')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': request.POST, 'fornecedores': fornecedores})
            
        # 3. Verificar se já existe produto com a mesma descrição e valor (se a descrição foi fornecida)
        if descricao and Produto.objects.filter(descricao=descricao, valor_unitario=valor_unitario).exists():
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com a mesma descrição e valor.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': request.POST, 'fornecedores': fornecedores})
        
        # Cria o produto
        produto = Produto(
            nome=nome,
            descricao=descricao,
            valor_unitario=valor_unitario,
            unidade_medida=unidade_medida,
            codigo=codigo,
            fornecedor=Fornecedor.objects.get(id=fornecedor_id) if fornecedor_id else None
        )
        produto.save()
        
        # Registra o log de inclusão
        from .models import ProdutoLog
        
        ProdutoLog.objects.create(
            produto=produto,
            nome_produto=produto.nome,
            acao='INCLUSAO',
            detalhes=f"Valor: R$ {produto.valor_unitario}, Unidade: {produto.unidade_medida}, Código: {produto.codigo or '-'}",
            usuario=request.user.username if request.user.is_authenticated else 'Anônimo',
            ip=get_client_ip(request)
        )
        
        messages.success(request, 'Produto adicionado com sucesso!')
        return redirect('produtos:lista')
    
    fornecedores = Fornecedor.objects.all()
    return render(request, 'produtos/form.html', {'fornecedores': fornecedores})

def editar(request, pk):
    """Edita um produto existente"""
    produto = get_object_or_404(Produto, pk=pk)
    nome_original = produto.nome
    valor_original = produto.valor_unitario
    codigo_original = produto.codigo or '-'
    ativo_original = produto.ativo
    
    if request.method == 'POST':
        # Extrai os dados do formulário
        produto.nome = request.POST.get('nome')
        produto.descricao = request.POST.get('descricao')
        produto.valor_unitario = request.POST.get('valor_unitario').replace(',', '.')
        produto.unidade_medida = request.POST.get('unidade_medida')
        produto.codigo = request.POST.get('codigo')
        produto.ativo = 'ativo' in request.POST
        fornecedor_id = request.POST.get('fornecedor')
        
        # Valida os dados
        if not produto.nome:
            messages.error(request, 'O nome do produto é obrigatório.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})
            
        # Remove espaços extras do nome para evitar duplicações por espaços
        produto.nome = produto.nome.strip()
        
        # Verificação de produto duplicado pelo nome (PRIORITÁRIA)
        nome_existente = Produto.objects.filter(nome__iexact=produto.nome).exclude(pk=pk).first()
        if nome_existente:
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o nome "{produto.nome}".')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})
        
        try:
            produto.valor_unitario = Decimal(produto.valor_unitario)
            if produto.valor_unitario <= 0:
                raise ValueError("Valor deve ser maior que zero")
        except:
            messages.error(request, 'O valor unitário deve ser um número positivo.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})
        
        # Verificações secundárias de duplicação
        # 2. Verificar se já existe produto com o mesmo código (se o código foi fornecido)
        if produto.codigo and Produto.objects.filter(codigo=produto.codigo).exclude(pk=pk).exists():
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o código "{produto.codigo}".')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})
            
        # 3. Verificar se já existe produto com a mesma descrição e valor (se a descrição foi fornecida)
        if produto.descricao and Produto.objects.filter(descricao=produto.descricao, valor_unitario=produto.valor_unitario).exclude(pk=pk).exists():
            messages.error(request, f'PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com a mesma descrição e valor.')
            fornecedores = Fornecedor.objects.all()
            return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})
        
        # Salva o produto
        produto.fornecedor = Fornecedor.objects.get(id=fornecedor_id) if fornecedor_id else None
        produto.save()
        
        # Registra o log de edição
        from .models import ProdutoLog
        
        alteracoes = []
        if nome_original != produto.nome:
            alteracoes.append(f"Nome alterado de '{nome_original}' para '{produto.nome}'")
        if valor_original != produto.valor_unitario:
            alteracoes.append(f"Valor alterado de R$ {valor_original} para R$ {produto.valor_unitario}")
        if codigo_original != (produto.codigo or '-'):
            alteracoes.append(f"Código alterado de '{codigo_original}' para '{produto.codigo or '-'}'")
        if ativo_original != produto.ativo:
            alteracoes.append(f"Status alterado de {'Ativo' if ativo_original else 'Inativo'} para {'Ativo' if produto.ativo else 'Inativo'}")
        
        detalhes = "; ".join(alteracoes) if alteracoes else "Nenhuma alteração significativa"
            
        ProdutoLog.objects.create(
            produto=produto,
            nome_produto=produto.nome,
            acao='EDICAO',
            detalhes=detalhes,
            usuario=request.user.username if request.user.is_authenticated else 'Anônimo',
            ip=get_client_ip(request)
        )
        
        messages.success(request, 'Produto atualizado com sucesso!')
        return redirect('produtos:lista')
    
    fornecedores = Fornecedor.objects.all()
    return render(request, 'produtos/form.html', {'produto': produto, 'fornecedores': fornecedores})

def importar(request):
    """Importa produtos de um arquivo Excel"""
    if request.method == 'POST' and 'arquivo' in request.FILES:
        arquivo = request.FILES['arquivo']
        
        # Verifica a extensão do arquivo
        if not arquivo.name.endswith(('.xlsx', '.xls')):
            messages.error(request, 'Formato de arquivo não suportado. Por favor, envie um arquivo Excel (.xlsx ou .xls).')
            return redirect('produtos:importar')
        
        try:
            # Carrega o arquivo Excel
            df = pd.read_excel(arquivo)
            
            # Verifica se as colunas necessárias existem
            colunas_necessarias = ['Nome do Produto', 'Valor Unitário']
            for coluna in colunas_necessarias:
                if coluna not in df.columns:
                    messages.error(request, f'Coluna "{coluna}" não encontrada no arquivo. Verifique o formato do arquivo.')
                    return redirect('produtos:importar')
            
            # Mapeia os nomes das colunas no Excel para os campos do modelo
            mapeamento = {
                'Nome do Produto': 'nome',
                'Descrição': 'descricao',
                'Valor Unitário': 'valor_unitario',
                'Unidade de Medida': 'unidade_medida',
                'Código': 'codigo',
                'Fornecedor (CNPJ)': 'fornecedor_cnpj',
            }
            
            produtos_importados = 0
            produtos_atualizados = 0
            produtos_ignorados = 0
            erros_importacao = []
            
            # Inicia uma transação para garantir consistência
            with transaction.atomic():
                for idx, row in df.iterrows():
                    dados_produto = {}
                    
                    # Extrai os dados da linha com base no mapeamento
                    for coluna_excel, campo_modelo in mapeamento.items():
                        if coluna_excel in df.columns and not pd.isna(row[coluna_excel]):
                            dados_produto[campo_modelo] = row[coluna_excel]
                    
                    # Buscar fornecedor pelo CNPJ, se informado
                    fornecedor_obj = None
                    if 'fornecedor_cnpj' in dados_produto:
                        fornecedor_cnpj = str(dados_produto['fornecedor_cnpj']).strip()
                        try:
                            fornecedor_obj = Fornecedor.objects.get(cnpj=fornecedor_cnpj)
                        except Fornecedor.DoesNotExist:
                            erros_importacao.append(f"Linha {idx+2}: Fornecedor com CNPJ '{fornecedor_cnpj}' não encontrado. Produto importado sem fornecedor.")
                        del dados_produto['fornecedor_cnpj']
                    
                    # Verifica campos obrigatórios
                    if 'nome' not in dados_produto or 'valor_unitario' not in dados_produto:
                        erros_importacao.append(f"Linha {idx+2}: Nome ou valor unitário não informados")
                        continue
                        
                    # Remove espaços extras do nome para evitar duplicações por espaços
                    if 'nome' in dados_produto:
                        dados_produto['nome'] = dados_produto['nome'].strip()
                    
                    # Verifica se o produto já existe (por código)
                    produto_existente = None
                    if 'codigo' in dados_produto and dados_produto['codigo']:
                        try:
                            produto_existente = Produto.objects.get(codigo=dados_produto['codigo'])
                        except Produto.DoesNotExist:
                            pass
                    
                    # Se não encontrou pelo código, verifica pelo nome
                    if not produto_existente:
                        try:
                            produto_existente = Produto.objects.get(nome__iexact=dados_produto['nome'])
                        except Produto.DoesNotExist:
                            pass
                        except Produto.MultipleObjectsReturned:
                            erros_importacao.append(f"Linha {idx+2}: Múltiplos produtos com o nome '{dados_produto['nome']}' encontrados")
                            continue
                    
                    # Se não encontrou pelo nome e há descrição e valor, verifica por esses critérios
                    if not produto_existente and 'descricao' in dados_produto:
                        try:
                            produto_existente = Produto.objects.get(
                                descricao=dados_produto['descricao'], 
                                valor_unitario=dados_produto['valor_unitario']
                            )
                        except Produto.DoesNotExist:
                            pass
                        except Produto.MultipleObjectsReturned:
                            erros_importacao.append(f"Linha {idx+2}: Múltiplos produtos com a mesma descrição e valor encontrados")
                            continue
                    
                    if produto_existente:
                        # Atualiza o produto existente
                        for campo, valor in dados_produto.items():
                            setattr(produto_existente, campo, valor)
                        produto_existente.fornecedor = fornecedor_obj
                        produto_existente.save()
                        produtos_atualizados += 1
                    else:
                        # Verifica se já existe produto com o mesmo nome (duplicado na importação)
                        if Produto.objects.filter(nome__iexact=dados_produto['nome']).exists():
                            erros_importacao.append(f"Linha {idx+2}: PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o nome '{dados_produto['nome']}'")
                            produtos_ignorados += 1
                            continue
                        
                        # Verifica se já existe produto com o mesmo código
                        if 'codigo' in dados_produto and dados_produto['codigo'] and Produto.objects.filter(codigo=dados_produto['codigo']).exists():
                            erros_importacao.append(f"Linha {idx+2}: PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com o código '{dados_produto['codigo']}'")
                            produtos_ignorados += 1
                            continue
                        
                        # Verifica se já existe produto com a mesma descrição e valor
                        if 'descricao' in dados_produto and Produto.objects.filter(
                            descricao=dados_produto['descricao'], 
                            valor_unitario=dados_produto['valor_unitario']
                        ).exists():
                            erros_importacao.append(f"Linha {idx+2}: PRODUTO NÃO PODE SER CADASTRADO DUPLICADO. Já existe um produto com a mesma descrição e valor")
                            produtos_ignorados += 1
                            continue
                        
                        # Cria um novo produto
                        # Define a unidade de medida padrão se não informada
                        if 'unidade_medida' not in dados_produto:
                            dados_produto['unidade_medida'] = 'Unidade'
                        
                        if fornecedor_obj:
                            dados_produto['fornecedor'] = fornecedor_obj
                        
                        novo_produto = Produto.objects.create(**dados_produto)
                        produtos_importados += 1
                        
                        # Registra o log de importação
                        from .models import ProdutoLog
                        
                        ProdutoLog.objects.create(
                            produto=novo_produto,
                            nome_produto=novo_produto.nome,
                            acao='IMPORTACAO',
                            detalhes=f"Produto importado. Valor: R$ {novo_produto.valor_unitario}, Código: {novo_produto.codigo or '-'}",
                            usuario=request.user.username if request.user.is_authenticated else 'Anônimo',
                            ip=get_client_ip(request)
                        )
            
            # Mensagens de sucesso e erro
            if produtos_importados > 0 or produtos_atualizados > 0:
                messages.success(
                    request, 
                    f'Importação concluída: {produtos_importados} produtos importados, {produtos_atualizados} produtos atualizados. '
                    f'{produtos_ignorados} PRODUTOS NÃO CADASTRADOS POR SEREM DUPLICADOS.'
                )
            
            if erros_importacao:
                messages.warning(
                    request, 
                    f'Atenção: {len(erros_importacao)} erros encontrados. Os primeiros 5 erros são: ' + 
                    ', '.join(erros_importacao[:5]) + 
                    ('. E mais...' if len(erros_importacao) > 5 else '.')
                )
                
            return redirect('produtos:lista')
            
        except Exception as e:
            messages.error(request, f'Erro ao processar o arquivo: {str(e)}')
            return redirect('produtos:importar')
    
    return render(request, 'produtos/importar.html')

def download_modelo(request):
    """Download do modelo Excel para importação de produtos"""
    # Cria uma nova planilha
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Modelo Importação Produtos"
    
    # Define os cabeçalhos
    cabecalhos = ["Nome do Produto", "Descrição", "Valor Unitário", "Unidade de Medida", "Código", "Fornecedor (CNPJ)"]
    for col_num, cabecalho in enumerate(cabecalhos, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = cabecalho
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    # Adiciona exemplos
    exemplos = [
        ["Caderno Espiral", "Caderno espiral 200 folhas, capa dura", 15.90, "Unidade", "CAD001", "12345678000199"],
        ["Lápis HB", "Caixa com 12 lápis pretos HB", 12.50, "Caixa", "LAP002", ""],
    ]
    
    for row_num, exemplo in enumerate(exemplos, 2):
        for col_num, valor in enumerate(exemplo, 1):
            sheet.cell(row=row_num, column=col_num).value = valor
    
    # Ajusta a largura das colunas
    for col in sheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column].width = adjusted_width
    
    # Cria a resposta HTTP com o arquivo
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename=modelo_produtos.xlsx'
    
    # Salva o workbook na resposta
    workbook.save(response)
    
    return response

def excluir(request, pk):
    """Exclui um produto se ele não estiver relacionado a nenhum pedido"""
    try:
        produto = get_object_or_404(Produto, pk=pk)
        nome_produto = produto.nome
        codigo_produto = produto.codigo or '-'
        valor_produto = produto.valor_unitario
        
        # Verifica se o produto está relacionado a algum pedido
        from pedidos.models import ItemPedido
        
        pedidos_relacionados = ItemPedido.objects.filter(produto=produto)
        
        if pedidos_relacionados.exists():
            num_pedidos = pedidos_relacionados.values('pedido').distinct().count()
            messages.error(
                request, 
                f'PRODUTO NÃO PODE SER EXCLUÍDO. O produto "{produto.nome}" está relacionado a {num_pedidos} pedido(s).'
            )
        else:
            try:
                # Registra o log de exclusão antes de excluir o produto
                from .models import ProdutoLog
                
                ProdutoLog.objects.create(
                    produto=None,  # Produto será None pois será excluído
                    nome_produto=nome_produto,
                    acao='EXCLUSAO',
                    detalhes=f"Produto excluído. Valor: R$ {valor_produto}, Código: {codigo_produto}",
                    usuario=request.user.username if request.user.is_authenticated else 'Anônimo',
                    ip=get_client_ip(request)
                )
                
                produto.delete()
                messages.success(request, f'Produto "{nome_produto}" excluído com sucesso!')
            except Exception as e:
                messages.error(request, f'Erro ao excluir o produto: {str(e)}')
    except Exception as e:
        messages.error(request, f'Erro ao processar a solicitação de exclusão: {str(e)}')
    
    # Sempre redireciona para a lista, independentemente do resultado
    return redirect('produtos:lista')

def logs(request):
    """Exibe os logs de ações realizadas nos produtos"""
    from .models import ProdutoLog
    from django.utils import timezone
    from datetime import timedelta
    
    # Filtros
    acao = request.GET.get('acao', '')
    usuario = request.GET.get('usuario', '')
    data_inicio = request.GET.get('data_inicio', '')
    data_fim = request.GET.get('data_fim', '')
    produto = request.GET.get('produto', '')
    
    # Inicia com todos os logs
    logs = ProdutoLog.objects.all()
    
    # Aplica os filtros se forem fornecidos
    if acao:
        logs = logs.filter(acao=acao)
    
    if usuario:
        logs = logs.filter(usuario__icontains=usuario)
    
    if produto:
        logs = logs.filter(nome_produto__icontains=produto)
    
    if data_inicio:
        try:
            data_inicio_obj = timezone.datetime.strptime(data_inicio, '%Y-%m-%d').date()
            logs = logs.filter(data_hora__date__gte=data_inicio_obj)
        except:
            pass
    
    if data_fim:
        try:
            data_fim_obj = timezone.datetime.strptime(data_fim, '%Y-%m-%d').date()
            logs = logs.filter(data_hora__date__lte=data_fim_obj)
        except:
            pass
    
    # Se não houver filtros, limita aos últimos 30 dias por padrão
    if not any([acao, usuario, produto, data_inicio, data_fim]):
        data_limite = timezone.now() - timedelta(days=30)
        logs = logs.filter(data_hora__gte=data_limite)
    
    # Ordenação mais recente primeiro
    logs = logs.order_by('-data_hora')
    
    # Obtém lista de usuários para o filtro
    usuarios_unicos = ProdutoLog.objects.values_list('usuario', flat=True).distinct()
    
    context = {
        'logs': logs,
        'usuarios': usuarios_unicos,
        'filtros': {
            'acao': acao,
            'usuario': usuario,
            'produto': produto,
            'data_inicio': data_inicio,
            'data_fim': data_fim,
        }
    }
    
    return render(request, 'produtos/logs.html', context)

def excluir_todos(request):
    """Exibe página de confirmação e exclui todos os produtos do sistema"""
    from .models import Produto, ProdutoLog
    from django.db import connection
    import os
    import json
    from datetime import datetime
    from django.conf import settings
    
    total_produtos = Produto.objects.count()
    
    # Se não há produtos, redireciona para a lista com mensagem
    if total_produtos == 0:
        messages.warning(request, 'Não há produtos cadastrados para excluir.')
        return redirect('produtos:lista')
    
    # Se for uma solicitação POST, executa a exclusão
    if request.method == 'POST' and 'confirmar_exclusao' in request.POST:
        try:
            # Verifica dependências com pedidos
            try:
                from pedidos.models import ItemPedido
                pedidos_relacionados = ItemPedido.objects.all().count()
                if pedidos_relacionados > 0:
                    messages.error(
                        request, 
                        f'OPERAÇÃO CANCELADA. Existem produtos relacionados a pedidos. '
                        f'Exclua os pedidos primeiro antes de excluir os produtos.'
                    )
                    return redirect('produtos:lista')
            except Exception as e:
                messages.warning(request, f'Não foi possível verificar a existência de pedidos relacionados: {str(e)}')
            
            # Criar diretório de backup se não existir
            backup_dir = os.path.join(settings.BASE_DIR, 'backups')
            os.makedirs(backup_dir, exist_ok=True)
            
            # Realizar backup dos produtos antes de excluí-los
            produtos_data = list(Produto.objects.values())
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            backup_file = os.path.join(backup_dir, f'produtos_backup_{timestamp}.json')
            
            with open(backup_file, 'w', encoding='utf-8') as f:
                json.dump(produtos_data, f, ensure_ascii=False, indent=4)
            
            # Registra o log de exclusão para cada produto
            for produto in Produto.objects.all():
                try:
                    ProdutoLog.objects.create(
                        produto=None,
                        nome_produto=produto.nome,
                        acao='EXCLUSAO',
                        detalhes=f"Produto excluído em exclusão em massa. Valor: R$ {produto.valor_unitario}, Código: {produto.codigo or '-'}",
                        usuario=request.user.username if request.user.is_authenticated else 'Anônimo',
                        ip=get_client_ip(request)
                    )
                except Exception as e:
                    messages.warning(request, f'Erro ao registrar log para produto {produto.nome}: {str(e)}')
            
            # Exclui todos os produtos
            Produto.objects.all().delete()
            
            # Reseta a sequência de IDs para começar do 1 novamente
            with connection.cursor() as cursor:
                if 'sqlite' in connection.vendor:
                    cursor.execute("DELETE FROM sqlite_sequence WHERE name='produtos_produto';")
                elif 'postgresql' in connection.vendor:
                    cursor.execute("ALTER SEQUENCE produtos_produto_id_seq RESTART WITH 1;")
                elif 'mysql' in connection.vendor:
                    cursor.execute("ALTER TABLE produtos_produto AUTO_INCREMENT = 1;")
            
            messages.success(
                request, 
                f'Todos os {total_produtos} produtos foram excluídos com sucesso! '
                f'Um backup foi salvo em: {backup_file}'
            )
            
        except Exception as e:
            messages.error(request, f'Erro ao excluir produtos: {str(e)}')
        
        return redirect('produtos:lista')
    
    # Apenas exibe a página de confirmação
    return render(request, 'produtos/excluir_todos.html', {'total_produtos': total_produtos})

@login_required
def listar_fornecedores(request):
    fornecedores = Fornecedor.objects.all()
    fornecedores_ativos = fornecedores.filter(ativo=True).count()
    fornecedores_inativos = fornecedores.filter(ativo=False).count()
    cidades_diferentes = fornecedores.values('cidade').distinct().count()
    return render(request, 'produtos/fornecedores/listar.html', {
        'fornecedores': fornecedores,
        'fornecedores_ativos': fornecedores_ativos,
        'fornecedores_inativos': fornecedores_inativos,
        'cidades_diferentes': cidades_diferentes
    })

@login_required
def criar_fornecedor(request):
    if request.method == 'POST':
        # Campos obrigatórios
        razao_social = request.POST.get('razao_social')
        nome = request.POST.get('nome')
        cnpj = request.POST.get('cnpj')
        
        # Validação dos campos obrigatórios
        if not all([razao_social, cnpj]):
            messages.error(request, 'Os campos Razão Social e CNPJ/CPF são obrigatórios.')
            return render(request, 'produtos/fornecedores/criar.html', {'fornecedor': request.POST})
        
        # Verifica se já existe fornecedor com o mesmo CNPJ
        if Fornecedor.objects.filter(cnpj=cnpj).exists():
            messages.error(request, 'Já existe um fornecedor cadastrado com este CNPJ/CPF.')
            return render(request, 'produtos/fornecedores/criar.html', {'fornecedor': request.POST})
        
        # Processar o dia do pedido recorrente
        dia_pedido_recorrente = request.POST.get('dia_pedido_recorrente')
        if dia_pedido_recorrente:
            try:
                dia_pedido_recorrente = int(dia_pedido_recorrente)
                if dia_pedido_recorrente < 1 or dia_pedido_recorrente > 31:
                    messages.warning(request, 'O dia do pedido recorrente deve estar entre 1 e 31. Valor ajustado para o limite mais próximo.')
                    dia_pedido_recorrente = min(max(dia_pedido_recorrente, 1), 31)
            except ValueError:
                dia_pedido_recorrente = None
                messages.warning(request, 'Valor inválido para o dia do pedido recorrente. Campo foi ignorado.')
        else:
            dia_pedido_recorrente = None
        
        # Criação do fornecedor com todos os campos
        fornecedor = Fornecedor(
            razao_social=razao_social,
            nome=nome,
            cnpj=cnpj,
            inscricao_estadual=request.POST.get('inscricao_estadual'),
            endereco=request.POST.get('endereco'),
            cep=request.POST.get('cep'),
            cidade=request.POST.get('cidade'),
            estado=request.POST.get('estado'),
            telefone=request.POST.get('telefone'),
            telefone_secundario=request.POST.get('telefone_secundario'),
            email=request.POST.get('email'),
            email_secundario=request.POST.get('email_secundario'),
            website=request.POST.get('website'),
            pessoa_contato=request.POST.get('pessoa_contato'),
            cargo_contato=request.POST.get('cargo_contato'),
            descricao=request.POST.get('descricao'),
            dia_pedido_recorrente=dia_pedido_recorrente,
            tipo_fornecedor=request.POST.get('tipo_fornecedor', 'FIXO')
        )
        fornecedor.save()
        messages.success(request, 'Fornecedor criado com sucesso!')
        return redirect('fornecedores:listar_fornecedores')
    
    return render(request, 'produtos/fornecedores/criar.html')

@login_required
def editar_fornecedor(request, id):
    fornecedor = get_object_or_404(Fornecedor, id=id)
    
    if request.method == 'POST':
        # Campos obrigatórios
        razao_social = request.POST.get('razao_social')
        nome = request.POST.get('nome')
        cnpj = request.POST.get('cnpj')
        
        # Validação dos campos obrigatórios
        if not all([razao_social, cnpj]):
            messages.error(request, 'Os campos Razão Social e CNPJ/CPF são obrigatórios.')
            return render(request, 'produtos/fornecedores/criar.html', {'fornecedor': request.POST})
        
        # Processar o dia do pedido recorrente
        dia_pedido_recorrente = request.POST.get('dia_pedido_recorrente')
        if dia_pedido_recorrente:
            try:
                dia_pedido_recorrente = int(dia_pedido_recorrente)
                if dia_pedido_recorrente < 1 or dia_pedido_recorrente > 31:
                    messages.warning(request, 'O dia do pedido recorrente deve estar entre 1 e 31. Valor ajustado para o limite mais próximo.')
                    dia_pedido_recorrente = min(max(dia_pedido_recorrente, 1), 31)
            except ValueError:
                dia_pedido_recorrente = None
                messages.warning(request, 'Valor inválido para o dia do pedido recorrente. Campo foi ignorado.')
        else:
            dia_pedido_recorrente = None
        
        # Atualização de todos os campos
        fornecedor.razao_social = razao_social
        fornecedor.nome = nome
        fornecedor.cnpj = cnpj
        fornecedor.inscricao_estadual = request.POST.get('inscricao_estadual')
        fornecedor.endereco = request.POST.get('endereco')
        fornecedor.cep = request.POST.get('cep')
        fornecedor.cidade = request.POST.get('cidade')
        fornecedor.estado = request.POST.get('estado')
        fornecedor.telefone = request.POST.get('telefone')
        fornecedor.telefone_secundario = request.POST.get('telefone_secundario')
        fornecedor.email = request.POST.get('email')
        fornecedor.email_secundario = request.POST.get('email_secundario')
        fornecedor.website = request.POST.get('website')
        fornecedor.pessoa_contato = request.POST.get('pessoa_contato')
        fornecedor.cargo_contato = request.POST.get('cargo_contato')
        fornecedor.descricao = request.POST.get('descricao')
        fornecedor.dia_pedido_recorrente = dia_pedido_recorrente
        fornecedor.tipo_fornecedor = request.POST.get('tipo_fornecedor', 'FIXO')
        
        fornecedor.save()
        messages.success(request, 'Fornecedor atualizado com sucesso!')
        return redirect('fornecedores:listar_fornecedores')
    
    return render(request, 'produtos/fornecedores/criar.html', {'fornecedor': fornecedor})

@login_required
def excluir_fornecedor(request, id):
    fornecedor = get_object_or_404(Fornecedor, id=id)
    fornecedor.delete()
    messages.success(request, 'Fornecedor excluído com sucesso!')
    return redirect('fornecedores:listar_fornecedores')
